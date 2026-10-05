import re
from pathlib import Path
import pymupdf as fitz
from pptx import Presentation


def _clean_slide_text(text: str) -> str:
    """Filter out boilerplate course metadata (instructor affiliations, slide counters, copyright)."""
    lines = []
    for line in text.splitlines():
        l_strip = line.strip()
        if not l_strip:
            continue
        # Skip pure page numbers or slide counts
        if re.match(r"^(?:slide|page)?\s*\d+\s*(?:/\s*\d+)?$", l_strip, re.I):
            continue
        # Skip common academic boilerplate
        if re.search(r"\b(all rights reserved|copyright|nptel|coursera|dept of|department of|professor|prof\.|iit |university)\b", l_strip, re.I) and len(l_strip) < 60:
            continue
        lines.append(l_strip)
    return "\n".join(lines).strip()


_ocr_engine = None

def _get_ocr_engine():
    global _ocr_engine
    if _ocr_engine is None:
        try:
            from rapidocr_onnxruntime import RapidOCR
            _ocr_engine = RapidOCR()
        except Exception as e:
            print(f"[extract] Could not initialize RapidOCR: {e}")
            _ocr_engine = False
    return _ocr_engine if _ocr_engine is not False else None


def extract_document(path: str, output_slides_dir: str = None):
    """Return a list of slide/page dictionaries from PPTX, PDF, or TXT, and optionally save slide images."""
    p = Path(path)
    suffix = p.suffix.lower()
    slides_dir = Path(output_slides_dir) if output_slides_dir else None
    if slides_dir:
        slides_dir.mkdir(parents=True, exist_ok=True)

    if suffix == ".pptx":
        prs = Presentation(path)
        pages = []
        for i, slide in enumerate(prs.slides, 1):
            parts = []
            for shape in slide.shapes:
                if getattr(shape, "has_text_frame", False):
                    text = shape.text_frame.text.strip()
                    if text:
                        parts.append(text)
            clean_content = _clean_slide_text("\n".join(parts))
            pages.append({
                "number": i,
                "text": clean_content or "\n".join(parts).strip(),
                "raw_text": "\n".join(parts).strip(),
                "image_path": None
            })
        return pages

    if suffix == ".pdf":
        doc = fitz.open(path)
        pages = []
        ocr_engine = _get_ocr_engine()
        ocr_count = 0
        max_ocr_pages = 25  # Limit OCR to top 25 pages for speed while ensuring high syllabus depth

        for i, page in enumerate(doc, 1):
            raw_text = page.get_text("text").strip()

            # Scanned or image-based slide fallback to OCR
            if len(raw_text) < 15 and ocr_engine and ocr_count < max_ocr_pages:
                try:
                    pix = page.get_pixmap(dpi=150)
                    ocr_res, _ = ocr_engine(pix.tobytes("png"))
                    if ocr_res:
                        ocr_text = " ".join([line[1] for line in ocr_res]).strip()
                        if ocr_text:
                            raw_text = ocr_text
                            ocr_count += 1
                except Exception as ocr_err:
                    print(f"[extract] OCR error on page {i}: {ocr_err}")

            clean_content = _clean_slide_text(raw_text)
            image_path = None
            has_diagram = False

            # Check for authentic visual diagrams (raster images or vector drawings)
            images = page.get_images(full=True)
            drawings = page.get_drawings()
            # If the page has images or non-trivial vector drawings, it likely contains a diagram
            if len(images) > 0 or len(drawings) >= 6:
                has_diagram = True

            if slides_dir and has_diagram:
                slide_file = slides_dir / f"diagram_{i}.png"
                try:
                    pix = page.get_pixmap(dpi=150)
                    pix.save(str(slide_file))
                    image_path = str(slide_file)
                except Exception as img_err:
                    print(f"[extract] Could not save diagram image for page {i}: {img_err}")

            pages.append({
                "number": i,
                "text": clean_content or raw_text,
                "raw_text": raw_text,
                "image_path": image_path,
                "has_diagram": has_diagram
            })
        doc.close()
        return pages

    if suffix == ".txt":
        text = Path(path).read_text(encoding="utf-8", errors="ignore")
        chunks = [x.strip() for x in text.split("\f") if x.strip()]
        if not chunks:
            chunks = [text.strip()]
        return [{"number": i, "text": chunk, "raw_text": chunk, "image_path": None} for i, chunk in enumerate(chunks, 1)]

    raise ValueError("Only PPTX, PDF and TXT files are supported.")
